#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
复检：配色合规 / 页脚连续 / 越界 / 溢出 / 字体统一。

用法:
    python verify.py <pptx路径> [--palette palette.json]

palette.json: {"allowed": ["2563EB","14B8A6","F59E0B", ...]}
不传时用内置的宽松色板（覆盖常见设计系统色族 + 中性灰 + 语义红绿），
目的是抓"明显跑偏的杂色"，不是卡死某个具体色值。
"""
import sys, os, io, re, json
from collections import Counter, defaultdict

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
EMU_PX = 9525.0

# 内置宽松色板：tailwind 常见设计色族 + 中性 + 语义色
BUILTIN = {
    # 蓝 / 靛 / sky
    "2563EB", "1D4ED8", "1E40AF", "1E3A8A", "3B82F6", "60A5FA", "93C5FD", "BFDBFE",
    "6366F1", "4F46E5", "4338CA", "3730A3", "312E81", "818CF8", "A5B4FC", "C7D2FE", "E0E7FF",
    "0EA5E9", "0284C7", "0369A1", "38BDF8", "7DD3FC", "BAE6FD", "E0F2FE",
    # 青 / teal / cyan
    "14B8A6", "0D9488", "0F766E", "115E59", "134E4A", "2DD4BF", "5EEAD4", "99F6E4", "CCFBF1", "F0FDFA",
    "06B6D4", "0891B2", "0E7490", "155E75", "22D3EE", "67E8F9", "A5F3FC", "CFFAFE", "ECFEFF",
    # 琥珀 / 橙 / 黄
    "F59E0B", "D97706", "B45309", "92400E", "78350F", "FBBF24", "FCD34D", "FDE68A", "FEF3C7", "FFFBEB",
    "F97316", "FB923C", "EA580C", "C2410C", "9A3412", "FDBA74", "FED7AA", "FFEDD5",
    "EAB308", "CA8A04", "A16207", "FACC15", "FDE047", "FEF08A", "FEF9C3",
    # 绿 / emerald
    "10B981", "059669", "047857", "065F46", "064E3B", "34D399", "6EE7B7", "A7F3D0", "D1FAE5", "ECFDF5",
    "22C55E", "16A34A", "15803D", "4ADE80", "86EFAC", "BBF7D0", "DCFCE7", "F0FDF4",
    # 语义红（警示 / 错误示范）
    "EF4444", "DC2626", "B91C1C", "991B1B", "F87171", "FCA5A5", "FECACA", "FEF2F2",
    # 中性灰阶
    "FFFFFF", "000000", "0F172A", "1E293B", "334155", "475569", "64748B",
    "94A3B8", "CBD5E1", "E2E8F0", "F1F5F9", "F8FAFC",
    # 浅底蓝
    "EFF6FF", "DBEAFE",
}


def char_w(c):
    return 1.0 if ord(c) > 0x2E80 else 0.55


def estimate(tf, w_px):
    ins_h = 0.0
    try:
        ins_h = (tf.margin_left + tf.margin_right) / EMU_PX
    except Exception:
        pass
    avail = max(w_px - ins_h, 10)
    need = 0.0
    for p in tf.paragraphs:
        t = "".join(r.text for r in p.runs)
        if not t:
            need += 14
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        fs = max(sizes) if sizes else 18.0
        cw = sum(char_w(c) for c in t) * fs * 0.98
        lines = max(1, int(cw / avail) + (1 if cw % avail > 0.5 else 0))
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
        need += lines * fs * (ls or 1.15) * 1.22
    ins_v = 0.0
    try:
        ins_v = (tf.margin_top + tf.margin_bottom) / EMU_PX
    except Exception:
        pass
    return need + ins_v


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    path = sys.argv[1]
    allowed = BUILTIN
    if "--palette" in sys.argv:
        i = sys.argv.index("--palette")
        allowed = set(c.upper().lstrip("#")
                      for c in json.load(open(sys.argv[i + 1], encoding="utf-8"))["allowed"])

    from pptx import Presentation

    prs = Presentation(path)
    PW, PH = prs.slide_width / EMU_PX, prs.slide_height / EMU_PX

    bad = defaultdict(Counter)
    oob, overflow, footers = [], [], {}
    fonts = Counter()

    def fill_hex(sh):
        try:
            f = sh.fill
            if f.type is not None and int(f.type) == 1:
                return str(f.fore_color.rgb).upper()
        except Exception:
            pass
        return None

    def walk(shapes, sidx):
        for sh in shapes:
            fh = fill_hex(sh)
            if fh and fh not in allowed:
                bad[sidx][fh] += 1
            try:
                x, y = sh.left / EMU_PX, sh.top / EMU_PX
                w, h = sh.width / EMU_PX, sh.height / EMU_PX
            except Exception:
                x = y = w = h = 0
            if x < -2 or y < -2 or x + w > PW + 2 or y + h > PH + 2:
                oob.append((sidx, round(x), round(y), round(w), round(h)))
            if sh.has_text_frame:
                tf = sh.text_frame
                txt = "\n".join(p.text for p in tf.paragraphs)
                need = estimate(tf, w)
                if h > 0 and need > h * 1.15 and len(txt) > 4:
                    overflow.append((sidx, round(need), round(h), txt[:60]))
                if re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", txt.strip()):
                    footers[sidx] = txt.strip()
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            fonts[r.font.name] += 1
            if hasattr(sh, "shapes"):
                try:
                    walk(sh.shapes, sidx)
                except Exception:
                    pass

    n = 0
    for i, s in enumerate(prs.slides):
        walk(s.shapes, i)
        n = i + 1

    total_bad = sum(bad[i].total() for i in bad)
    print(f"文件: {path}")
    print(f"页面尺寸: {PW:.0f} x {PH:.0f} px    页数: {n}\n")

    print(f"=== 跑出色板的填充色: {total_bad} 处 ===")
    if not total_bad:
        print("  OK 无")
    else:
        tot = Counter()
        for i in bad:
            tot.update(bad[i])
        for c, v in tot.most_common(15):
            pg = sorted([i for i in bad if c in bad[i]])
            print(f"  #{c} x{v}  页{pg[:12]}")

    print(f"\n=== 页脚校验（应 = 页码 / {n}）===")
    bad_ft = []
    for i in sorted(footers):
        m = re.match(r"(\d+)\s*/\s*(\d+)", footers[i])
        num, den = int(m.group(1)), int(m.group(2))
        if num != i + 1 or den != n:
            bad_ft.append((i, footers[i], f"应为 {i + 1} / {n}"))
    if not bad_ft:
        print(f"  OK {len(footers)} 个页脚全部正确")
    else:
        for b in bad_ft[:25]:
            print(f"  X P{b[0]:02d} '{b[1]}'  {b[2]}")
        print(f"  共 {len(bad_ft)} 个不正确")

    print(f"\n=== 越界元素: {len(oob)} ===")
    for o in oob[:10]:
        print(f"  P{o[0]:02d} pos=({o[1]},{o[2]}) {o[3]}x{o[4]}")

    print(f"\n=== 文字溢出: {len(overflow)} ===")
    for o in overflow[:10]:
        print(f"  P{o[0]:02d} 需{o[1]}px 实{o[2]}px  '{o[3]}'")

    print("\n=== 字体 ===")
    for k, v in fonts.most_common(5):
        print(f"  {k} x{v}")


if __name__ == "__main__":
    main()
