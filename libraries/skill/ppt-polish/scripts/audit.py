#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
PPT 全量体检：越界 / 溢出 / 配色 / 字体 / 字号 / 页脚 / 时长文案。

用法:
    python audit.py <pptx路径> [--json 输出文件]

输出一份"这文件到底有什么毛病"的清单，作为后续精修的依据。
"""
import sys, os, io, re, json
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
EMU_PX = 9525.0


def fill_hex(sh):
    try:
        f = sh.fill
        if f.type is not None and int(f.type) == 1:
            return str(f.fore_color.rgb).upper()
    except Exception:
        pass
    return None


def font_hex(run):
    try:
        if run.font.color and run.font.color.rgb is not None:
            return str(run.font.color.rgb).upper()
    except Exception:
        pass
    return None


def char_w(c):
    """全角 1.0，半角 0.55"""
    return 1.0 if ord(c) > 0x2E80 else 0.55


def estimate(tf, w_px, h_px):
    """粗估文本框需要的高度（px）"""
    ins_h = 0.0
    try:
        ins_h = (tf.margin_left + tf.margin_right) / EMU_PX
    except Exception:
        pass
    avail = max(w_px - ins_h, 10)
    need = 0.0
    for p in tf.paragraphs:
        t = "".join(r.text for r in p.runs)
        if not t:
            need += 14
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        fs = max(sizes) if sizes else 18.0
        cw = sum(char_w(c) for c in t) * fs * 0.98
        lines = max(1, int(cw / avail) + (1 if cw % avail > 0.5 else 0))
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
        need += lines * fs * (ls or 1.15) * 1.22
    ins_v = 0.0
    try:
        ins_v = (tf.margin_top + tf.margin_bottom) / EMU_PX
    except Exception:
        pass
    return need + ins_v


def run(path):
    from pptx import Presentation

    prs = Presentation(path)
    PW, PH = prs.slide_width / EMU_PX, prs.slide_height / EMU_PX

    fills, fonts, sizes, fcolors = Counter(), Counter(), Counter(), Counter()
    oob, overflow, footers, times = [], [], [], []

    def walk(shapes, sidx):
        for sh in shapes:
            fh = fill_hex(sh)
            if fh:
                fills[fh] += 1
            try:
                x, y = sh.left / EMU_PX, sh.top / EMU_PX
                w, h = sh.width / EMU_PX, sh.height / EMU_PX
            except Exception:
                x = y = w = h = 0
            if x < -2 or y < -2 or x + w > PW + 2 or y + h > PH + 2:
                oob.append({"page": sidx, "x": round(x), "y": round(y),
                            "w": round(w), "h": round(h),
                            "text": (sh.text_frame.text[:50] if sh.has_text_frame else "")})
            if sh.has_text_frame:
                tf = sh.text_frame
                txt = "\n".join(p.text for p in tf.paragraphs)
                need = estimate(tf, w, h)
                if h > 0 and need > h * 1.15 and len(txt) > 4:
                    overflow.append({"page": sidx, "need": round(need),
                                     "have": round(h), "text": txt[:80]})
                if re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", txt.strip()):
                    footers.append({"page": sidx, "text": txt.strip()})
                if re.search(r"(分钟|小时|min\b)", txt):
                    times.append({"page": sidx, "text": txt[:70].replace("\n", " | ")})
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            fonts[r.font.name] += 1
                        if r.font.size:
                            sizes[round(r.font.size.pt, 1)] += 1
                        fc = font_hex(r)
                        if fc:
                            fcolors[fc] += 1
            if hasattr(sh, "shapes"):
                try:
                    walk(sh.shapes, sidx)
                except Exception:
                    pass

    n = 0
    for i, s in enumerate(prs.slides):
        walk(s.shapes, i)
        n = i + 1

    print(f"文件: {path}")
    print(f"页面尺寸: {PW:.0f} x {PH:.0f} px    页数: {n}\n")

    print("=== 填充色 TOP20 ===")
    for k, v in fills.most_common(20):
        print(f"  #{k}  x{v}")

    print("\n=== 字体 TOP8 ===")
    for k, v in fonts.most_common(8):
        print(f"  {k}  x{v}")

    print("\n=== 字号 TOP15 (pt, px=pt*4/3) ===")
    for k, v in sizes.most_common(15):
        print(f"  {k}pt (= {round(k * 4 / 3)}px)  x{v}")

    print("\n=== 文字色 TOP10 ===")
    for k, v in fcolors.most_common(10):
        print(f"  #{k}  x{v}")

    print(f"\n=== 越界元素: {len(oob)} ===")
    for o in oob[:20]:
        print(f"  P{o['page']:02d} pos=({o['x']},{o['y']}) {o['w']}x{o['h']}  '{o['text']}'")

    print(f"\n=== 文字溢出风险: {len(overflow)} ===")
    for o in overflow[:20]:
        print(f"  P{o['page']:02d} 需{o['need']}px 实{o['have']}px  '{o['text']}'")

    denoms = Counter()
    for f in footers:
        m = re.search(r"\d+\s*/\s*(\d+)", f["text"])
        if m:
            denoms[m.group(1)] += 1
    print(f"\n=== 页脚: {len(footers)} 个，分母分布 {dict(denoms)} ===")
    print(f"    （实际 {n} 页，分母应等于页数）")

    print(f"\n=== 时长文案: {len(times)} 处 ===")
    for t in times[:15]:
        print(f"  P{t['page']:02d} {t['text']}")

    if "--json" in sys.argv:
        idx = sys.argv.index("--json")
        if idx + 1 < len(sys.argv):
            with open(sys.argv[idx + 1], "w", encoding="utf-8") as f:
                json.dump({"page_px": [PW, PH], "slides": n,
                           "fills": fills.most_common(), "fonts": fonts.most_common(),
                           "sizes": sizes.most_common(), "fontcolors": fcolors.most_common(),
                           "oob": oob, "overflow": overflow,
                           "footers": footers, "times": times},
                          f, ensure_ascii=False, indent=1)
            print(f"\nJSON 已写入: {sys.argv[idx + 1]}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    run(sys.argv[1])
