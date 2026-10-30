#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
文本框重叠检测 —— 找出"文字压文字"的排版硬伤。

用法:
    python overlap.py <pptx路径> [阈值0.0-1.0，默认0.30]

规则:
  - 只对「两个都有非空文本」的框做两两比较
  - 完全包含（容器 / 背景）关系跳过
  - 交集面积 / 较小框面积 > 阈值 才算问题
"""
import sys, os, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
EMU_PX = 9525.0


def txt_of(sh):
    if sh.has_text_frame:
        return "".join(p.text for p in sh.text_frame.paragraphs).strip()
    return ""


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    path = sys.argv[1]
    thresh = float(sys.argv[2]) if len(sys.argv) > 2 else 0.30

    from pptx import Presentation

    prs = Presentation(path)

    def boxes(shapes, out):
        for sh in shapes:
            t = txt_of(sh)
            try:
                x, y = sh.left / EMU_PX, sh.top / EMU_PX
                w, h = sh.width / EMU_PX, sh.height / EMU_PX
            except Exception:
                x = y = w = h = 0
            if t and w > 0 and h > 0:
                out.append((x, y, w, h, t))
            if hasattr(sh, "shapes"):
                try:
                    boxes(sh.shapes, out)
                except Exception:
                    pass

    total = 0
    for i, s in enumerate(prs.slides):
        bs = []
        boxes(s.shapes, bs)
        hits = []
        for a in range(len(bs)):
            for b in range(a + 1, len(bs)):
                x1, y1, w1, h1, t1 = bs[a]
                x2, y2, w2, h2, t2 = bs[b]
                ix = max(0, min(x1 + w1, x2 + w2) - max(x1, x2))
                iy = max(0, min(y1 + h1, y2 + h2) - max(y1, y2))
                if ix <= 1 or iy <= 1:
                    continue
                contained = (
                    (x1 <= x2 + 1 and y1 <= y2 + 1 and
                     x1 + w1 >= x2 + w2 - 1 and y1 + h1 >= y2 + h2 - 1)
                    or
                    (x2 <= x1 + 1 and y2 <= y1 + 1 and
                     x2 + w2 >= x1 + w1 - 1 and y2 + h2 >= y1 + h1 - 1)
                )
                if contained:
                    continue
                ratio = (ix * iy) / min(w1 * h1, w2 * h2)
                if ratio > thresh:
                    hits.append((ratio, t1[:30], t2[:30],
                                 (round(x1), round(y1), round(w1), round(h1)),
                                 (round(x2), round(y2), round(w2), round(h2))))
        if hits:
            total += len(hits)
            print(f"\nP{i:02d}  {len(hits)} 处重叠")
            for h in sorted(hits, reverse=True)[:5]:
                print(f"   {h[0]:.0%}  '{h[1]}'  <->  '{h[2]}'")
                print(f"        A{h[3]}  B{h[4]}")

    print(f"\n>>> 合计 {total} 处文字框重叠（阈值 {thresh:.0%}）")


if __name__ == "__main__":
    main()
