#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
按规则批量精修 PPT（配色收敛 / 文字色 / 文案替换 / 页脚重编号），原地保存。

用法:
    python refine.py <pptx路径> rules.json [--dry]

rules.json 格式:
{
  "fill":  [{"from":"8B5CF6","to":"6366F1","pages":[1,7,21,33]}],
  "font":  [{"from":"7C3AED","to":"4338CA","pages":[34,35]}],
  "text":  [{"find":"60 分钟","replace":"约 2.5 小时","pages":[0,2]}],
  "footer": {"denominator":73, "offset_from":33, "offset":-1, "offset_to":71}
}

说明:
  - pages 省略 = 全局替换；pages 用 0-based 页索引数组
  - fill 会同时替换填充色和边框色
  - footer: 把 "NN / 旧分母" 重写为 "(NN+offset) / 新分母"，
            offset 只对 [offset_from, offset_to] 区间内的页生效
  - 改前请先备份；加 --dry 可只看会改多少处、不落盘
"""
import sys, os, io, re, json
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")


def build_index(rules, key):
    """{颜色: {页号集合 or None: 新色}} -> 返回 (全局映射, 按页映射)"""
    glob, bypage = {}, {}
    for r in rules.get(key, []):
        src = r["from"].upper().lstrip("#")
        dst = r["to"].upper().lstrip("#")
        if "pages" in r and r["pages"] is not None:
            bypage.setdefault(src, {})  # 占位
            for p in r["pages"]:
                bypage.setdefault(src, {}).setdefault(p, dst)
            bypage[src]["__pages__"] = set(r["pages"])
        else:
            glob[src] = dst
    return glob, bypage


def pick(src, pno, glob, bypage):
    if src in glob:
        return glob[src]
    if src in bypage:
        return bypage[src].get(pno)
    return None


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)

    path = sys.argv[1]
    rules = json.load(open(sys.argv[2], encoding="utf-8"))
    dry = "--dry" in sys.argv

    from pptx import Presentation
    from pptx.dml.color import RGBColor

    fill_g, fill_p = build_index(rules, "fill")
    font_g, font_p = build_index(rules, "font")
    text_rules = rules.get("text", [])
    footer = rules.get("footer")

    stat = Counter()
    prs = Presentation(path)

    def get_fill(sh):
        try:
            f = sh.fill
            if f.type is not None and int(f.type) == 1:
                return str(f.fore_color.rgb).upper()
        except Exception:
            pass
        return None

    def apply_line(sh, pno):
        try:
            ln = sh.line
            if ln.fill.type is not None and int(ln.fill.type) == 1:
                lc = str(ln.fill.fore_color.rgb).upper()
                dst = pick(lc, pno, fill_g, fill_p)
                if dst:
                    ln.fill.solid()
                    ln.fill.fore_color.rgb = RGBColor.from_string(dst)
                    stat[f"LINE P{pno:02d} #{lc}->#{dst}"] += 1
        except Exception:
            pass

    def walk(shapes, pno):
        for sh in shapes:
            fh = get_fill(sh)
            if fh:
                dst = pick(fh, pno, fill_g, fill_p)
                if dst:
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = RGBColor.from_string(dst)
                    stat[f"FILL P{pno:02d} #{fh}->#{dst}"] += 1
            apply_line(sh, pno)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        try:
                            if r.font.color and r.font.color.rgb is not None:
                                c = str(r.font.color.rgb).upper()
                                dst = pick(c, pno, font_g, font_p)
                                if dst:
                                    r.font.color.rgb = RGBColor.from_string(dst)
                                    stat[f"FONT P{pno:02d} #{c}->#{dst}"] += 1
                        except Exception:
                            pass
            if hasattr(sh, "shapes"):
                try:
                    walk(sh.shapes, pno)
                except Exception:
                    pass

    n = 0
    for i, slide in enumerate(prs.slides):
        walk(slide.shapes, i)
        n = i + 1

        # 文案替换
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for tr in text_rules:
                    if "pages" in tr and tr["pages"] is not None and i not in tr["pages"]:
                        continue
                    if tr["find"] not in p.text:
                        continue
                    runs = p.runs
                    if not runs:
                        continue
                    new_text = p.text.replace(tr["find"], tr["replace"])
                    runs[0].text = new_text
                    for r in runs[1:]:
                        r.text = ""
                    stat[f"TEXT P{i:02d} '{tr['find']}'->'{tr['replace']}'"] += 1

        # 页脚重编号
        if footer:
            den = footer.get("denominator")
            off = footer.get("offset", 0)
            ofrom = footer.get("offset_from", 0)
            oto = footer.get("offset_to", 10 ** 9)
            for sh in slide.shapes:
                if not sh.has_text_frame:
                    continue
                for p in sh.text_frame.paragraphs:
                    m = re.fullmatch(r"\s*(\d+)\s*/\s*(\d+)\s*", p.text)
                    if not m:
                        continue
                    num, old_den = int(m.group(1)), int(m.group(2))
                    if den is None:
                        den = old_den
                    if ofrom <= i <= oto:
                        num += off
                    new_text = f"{num} / {den}"
                    if p.text.strip() != new_text:
                        runs = p.runs
                        if runs:
                            runs[0].text = new_text
                            for r in runs[1:]:
                                r.text = ""
                            stat[f"FOOT P{i:02d} '{p.text.strip()}'->'{new_text}'"] += 1

    print(f"文件: {path}   页数: {n}")
    print(f"\n=== 变更统计（共 {sum(stat.values())} 处）===")
    for k, v in stat.most_common(50):
        print(f"  {k}  x{v}")
    if len(stat) > 50:
        print(f"  ... 另有 {len(stat) - 50} 类")

    if dry:
        print("\n[dry-run] 未写入磁盘")
        return

    prs.save(path)
    print(f"\n已保存: {path}")


if __name__ == "__main__":
    main()
